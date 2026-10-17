"""Load the exported tables, build the star schema, check it, and export it for Power BI.

    python -m gpstock.build              # uses data/source (run gpstock.generate first)
    python -m gpstock.build --generate   # make fresh synthetic data first
"""
import argparse
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

from gpstock import DATA_DIR
from gpstock.generate import END, START, generate

SOURCE_TABLES = ["sites", "categories", "items", "site_stock", "movements", "batches", "requests", "users"]
EXPORT = {
    "DimDate": "SELECT * FROM dim_date",
    "DimSite": "SELECT * FROM dim_site",
    "DimItem": "SELECT * FROM dim_item",
    "DimStaff": "SELECT * FROM dim_staff",
    "FactMovement": "SELECT * EXCLUDE (moved_at) FROM fact_movement",
    "FactDailyStock": "SELECT * FROM fact_daily_stock",
    "FactRequest": "SELECT * FROM fact_request",
    "FactBatch": "SELECT * FROM fact_batch",
    "ItemPerformance": "SELECT * FROM mart_item_performance",
}

# Each check returns the rows that break it.
CHECKS = {
    "Stock never goes negative": "SELECT * FROM fact_daily_stock WHERE closing_stock < 0",
    "Final stock matches the app's current stock": """
        SELECT * FROM src_site_stock s JOIN fact_daily_stock f
          ON f.site_id = s.site_id AND f.item_id = s.item_id AND f.date = (SELECT max(date) FROM dim_date)
        WHERE f.closing_stock <> s.quantity""",
    "Batch quantities add up to stock": """
        SELECT * FROM (SELECT site_id, item_id, sum(quantity_now) AS q FROM fact_batch GROUP BY ALL) b
        JOIN src_site_stock s USING (site_id, item_id) WHERE b.q <> s.quantity""",
    "Every movement has a known item and site": """
        SELECT * FROM fact_movement WHERE item_id NOT IN (SELECT item_id FROM dim_item)
           OR site_id NOT IN (SELECT site_id FROM dim_site)""",
    "Transfers out equal transfers in": """
        SELECT 1 FROM fact_movement HAVING sum(quantity) FILTER (WHERE movement_kind = 'transfer')
                                         <> sum(quantity) FILTER (WHERE movement_kind = 'transfer in')""",
}


def build(con: duckdb.DuckDBPyConnection, source=None) -> pd.DataFrame:
    source = Path(source or DATA_DIR / "source")
    for t in SOURCE_TABLES:
        con.execute(f"CREATE OR REPLACE TABLE src_{t} AS SELECT * FROM read_csv_auto(?, header = true)",
                    [str(source / f"{t}.csv")])
    sql = (Path(__file__).parent / "sql/model.sql").read_text()
    con.execute(sql.replace("$START", f"DATE '{START}'").replace("$END", f"DATE '{END}'"))
    return pd.DataFrame([{"check": name, "failing_rows": con.execute(f"SELECT count(*) FROM ({q})").fetchone()[0]}
                         for name, q in CHECKS.items()])


def export(con: duckdb.DuckDBPyConnection, out=None) -> dict:
    out = Path(out or DATA_DIR / "powerbi")
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for name, sql in EXPORT.items():
        df = con.execute(sql).df()
        df.to_csv(out / f"{name}.csv", index=False)
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in row])
        t = Table(displayName=name, ref=f"A1:{get_column_letter(len(df.columns))}{max(len(df), 1) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        counts[name] = len(df)
    wb.save(out / "GP_Stock_PowerBI.xlsx")
    return counts


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--generate", action="store_true")
    args = ap.parse_args(argv)
    if args.generate or not (DATA_DIR / "source/movements.csv").exists():
        generate()
    con = duckdb.connect(str(DATA_DIR / "gp_stock.duckdb"))
    checks = build(con)
    for r in checks.itertuples():
        print(f"  [{'pass' if r.failing_rows == 0 else 'FAIL'}] {r.check} ({r.failing_rows})")
    if checks.failing_rows.any():
        raise SystemExit("Data checks failed.")
    for name, n in export(con).items():
        print(f"  {name:<16} {n:>7,} rows")
    con.close()


if __name__ == "__main__":
    main()
